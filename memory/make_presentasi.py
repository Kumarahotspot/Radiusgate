from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

DARK = RGBColor(0x0B, 0x13, 0x20)
TEAL = RGBColor(0x0F, 0x76, 0x6E)
TEAL_LIGHT = RGBColor(0x14, 0xB8, 0xA6)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
SLATE = RGBColor(0x47, 0x55, 0x69)
LIGHT = RGBColor(0xF1, 0xF5, 0xF9)
AMBER = RGBColor(0xF5, 0x9E, 0x0B)

LOGO = "/app/frontend/public/logo.png"
LOGO_WHITE = "/app/frontend/public/logo-white.png"

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]
SW, SH = prs.slide_width, prs.slide_height


def add_slide(bg=WHITE):
    s = prs.slides.add_slide(BLANK)
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = bg
    return s


def box(s, x, y, w, h):
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tb.text_frame.word_wrap = True
    return tb


def para(tb, text, size=18, color=SLATE, bold=False, align=PP_ALIGN.LEFT, first=False, space_after=6):
    p = tb.text_frame.paragraphs[0] if first else tb.text_frame.add_paragraph()
    p.alignment = align
    p.space_after = Pt(space_after)
    r = p.add_run()
    r.text = text
    r.font.size = Pt(size)
    r.font.color.rgb = color
    r.font.bold = bold
    r.font.name = "Calibri"
    return p


def rect(s, x, y, w, h, fill, line=None):
    from pptx.enum.shapes import MSO_SHAPE
    sh = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    sh.adjustments[0] = 0.08
    sh.fill.solid()
    sh.fill.fore_color.rgb = fill
    if line:
        sh.line.color.rgb = line
        sh.line.width = Pt(1)
    else:
        sh.line.fill.background()
    sh.shadow.inherit = False
    return sh


def badge(s, x, y, text, w=2.6):
    b = rect(s, x, y, w, 0.42, TEAL)
    tf = b.text_frame
    tf.word_wrap = False
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = text
    r.font.size = Pt(12)
    r.font.bold = True
    r.font.color.rgb = WHITE
    return b


# ---------- Slide 1: Cover ----------
s = add_slide(DARK)
s.shapes.add_picture(LOGO_WHITE, Inches(5.42), Inches(1.1), height=Inches(1.5))
t = box(s, 1.2, 2.9, 10.9, 2.6)
para(t, "RadiusGate", 60, WHITE, True, PP_ALIGN.CENTER, True)
para(t, "Sistem Absensi Sekolah Modern — Wajah, QR, RFID & NIS dalam Satu Kiosk", 22, TEAL_LIGHT, False, PP_ALIGN.CENTER)
para(t, "Semua modul tersedia. Sekolah bebas menentukan metodenya sendiri.", 16, RGBColor(0x94, 0xA3, 0xB8), False, PP_ALIGN.CENTER)
badge(s, 5.17, 6.0, "Penawaran Program Pilot Sekolah", 3.0)

# ---------- Slide 2: Masalah ----------
s = add_slide()
badge(s, 0.8, 0.7, "MASALAH", 1.6)
t = box(s, 0.8, 1.3, 11.7, 1.0)
para(t, "Absensi Manual Memakan Waktu & Rawan Curang", 34, DARK, True, True)
problems = [
    ("Titip absen & manipulasi", "Absen kertas/cek manual mudah dipalsukan — tidak ada bukti kehadiran yang sah."),
    ("Rekap memakan waktu", "Guru piket menghitung kehadiran & keterlambatan manual setiap hari."),
    ("Orang tua tidak tahu", "Orang tua baru sadar anak tidak masuk setelah dipanggil sekolah."),
    ("Mesin absensi mahal", "Mesin fingerprint/RFID konvensional jutaan rupiah dan terkunci satu vendor."),
    ("Administrasi SPP terpisah", "Tagihan, kuitansi, dan pengingat pembayaran dikerjakan terpisah dari absensi."),
]
y = 2.5
for title, desc in problems:
    rect(s, 0.8, y, 11.7, 0.82, LIGHT)
    tb = box(s, 1.05, y + 0.08, 11.2, 0.7)
    para(tb, "✕  " + title, 15, RGBColor(0xDC, 0x26, 0x26), True, first=True)
    para(tb, desc, 12.5, SLATE)
    y += 0.95

# ---------- Slide 3: Solusi ----------
s = add_slide()
badge(s, 0.8, 0.7, "SOLUSI", 1.6)
t = box(s, 0.8, 1.3, 11.7, 1.0)
para(t, "RadiusGate: Kiosk Web + Portal Lengkap", 34, DARK, True, True)
t = box(s, 0.8, 2.3, 6.2, 4.6)
para(t, "Cukup tablet / HP bekas + browser", 20, TEAL, True, first=True)
for x in [
    "Kiosk berjalan di browser — tanpa beli mesin absensi mahal.",
    "Verifikasi wajah AI (ArcFace) + liveness anti foto/video palsu.",
    "GPS geofence: absen hanya sah di dalam area sekolah.",
    "Offline-first: internet down pun absen tetap jalan, sinkron otomatis.",
    "Suara sapaan otomatis & dwibahasa (Indonesia / Inggris).",
    "Full page otomatis via PWA — cocok untuk tablet di gerbang.",
]:
    para(t, "✓  " + x, 15, SLATE)
rect(s, 7.4, 2.3, 5.1, 4.3, DARK)
tb = box(s, 7.7, 2.6, 4.5, 3.8)
para(tb, "Satu akun, empat peran:", 16, TEAL_LIGHT, True, first=True)
para(tb, "Admin Sekolah — data, laporan, pengaturan", 13.5, WHITE)
para(tb, "Guru — absensi per mata pelajaran", 13.5, WHITE)
para(tb, "Karyawan — lembur & penggajian", 13.5, WHITE)
para(tb, "Orang Tua — pantau anak, izin/sakit, tagihan SPP", 13.5, WHITE)

# ---------- Slide 4: 4 Metode ----------
s = add_slide()
badge(s, 0.8, 0.7, "BEBAS PILIH", 1.9)
t = box(s, 0.8, 1.3, 11.7, 1.6)
para(t, "4 Metode Absensi dalam 1 Kiosk", 34, DARK, True, True)
para(t, "Sekolah menentukan sendiri metode yang dipakai — bisa satu, bisa semuanya sekaligus.", 15, SLATE)
methods = [
    ("Wajah (AI)", "ArcFace + liveness.\nAnti titip absen,\nverifikasi < 3 detik."),
    ("Kartu RFID / NFC", "Tap kartu di reader USB/OTG.\nTanpa sentuh layar.\nMode registrasi massal."),
    ("Kartu QR Code", "QR unik anti-palsu per orang,\nauto terdeteksi kamera.\nCetak massal per kelas (PDF)."),
    ("NIS / NIP Manual", "Keypad angka besar di layar,\nbunyi klik, foto profil\nkonfirmasi. Cadangan pasti."),
]
x = 0.8
for title, desc in methods:
    rect(s, x, 3.1, 2.95, 3.4, LIGHT, TEAL)
    tb = box(s, x + 0.18, 3.3, 2.6, 3.1)
    para(tb, title, 17, TEAL, True, PP_ALIGN.CENTER, True)
    for line in desc.split("\n"):
        para(tb, line, 12, SLATE, False, PP_ALIGN.CENTER)
    x += 3.13

# ---------- Slide 5: Semua modul ----------
s = add_slide()
badge(s, 0.8, 0.7, "SEMUA MODUL TERSEDIA", 2.9)
t = box(s, 0.8, 1.3, 11.7, 1.0)
para(t, "Lebih dari Sekadar Absensi", 34, DARK, True, True)
mods = [
    ("Absensi siswa & guru", "Harian real-time + per mata pelajaran oleh guru mapel."),
    ("Laporan & rekap", "Harian/bulanan per siswa & kelas, ekspor Excel & PDF."),
    ("SPP & penagihan", "Tagihan, cicilan, kuitansi PDF otomatis, pembayaran online."),
    ("Portal orang tua", "Pantau kehadiran anak, ajukan izin/sakit, lihat tagihan."),
    ("Notifikasi WhatsApp", "Otomatis ke orang tua setiap anak absen."),
    ("Lembur & penggajian", "Pengajuan lembur karyawan + perhitungan gaji otomatis."),
    ("Billing transparan", "Tagihan SaaS dihitung dari jumlah siswa aktif per bulan."),
    ("Screensaver kiosk", "Papan informasi sekolah tampil saat kiosk idle."),
]
y = 2.4
for i in range(0, len(mods), 2):
    for j, (title, desc) in enumerate(mods[i:i + 2]):
        x = 0.8 + j * 6.05
        rect(s, x, y, 5.85, 1.05, LIGHT)
        tb = box(s, x + 0.25, y + 0.12, 5.4, 0.85)
        para(tb, title, 15, TEAL, True, first=True)
        para(tb, desc, 12, SLATE)
    y += 1.25

# ---------- Slide 6: Cara kerja ----------
s = add_slide()
badge(s, 0.8, 0.7, "CARA KERJA", 1.9)
t = box(s, 0.8, 1.3, 11.7, 1.0)
para(t, "Terapkan dalam 3 Langkah", 34, DARK, True, True)
steps = [
    ("1", "Daftarkan Data", "Impor siswa/guru dari Excel/CSV, ambil sampel wajah, dan/atau registrasi kartu RFID & cetak kartu QR massal."),
    ("2", "Letakkan Tablet di Gerbang", "Buka kiosk di browser tablet/HP, pasang sebagai PWA agar full page — siswa absen dengan wajah, QR, kartu, atau NIS."),
    ("3", "Pantau Otomatis", "Dashboard admin, portal orang tua, notifikasi WhatsApp, dan laporan bulanan berjalan sendiri."),
]
x = 0.8
for n, title, desc in steps:
    rect(s, x, 2.6, 3.95, 3.6, DARK)
    tb = box(s, x + 0.3, 2.85, 3.35, 3.2)
    para(tb, n, 40, TEAL_LIGHT, True, PP_ALIGN.LEFT, True)
    para(tb, title, 19, WHITE, True)
    para(tb, desc, 13, RGBColor(0xCB, 0xD5, 0xE1))
    x += 4.15

# ---------- Slide 7: Penawaran pilot ----------
s = add_slide(DARK)
s.shapes.add_picture(LOGO_WHITE, Inches(0.8), Inches(0.7), height=Inches(0.9))
t = box(s, 0.8, 1.9, 11.7, 1.4)
para(t, "Penawaran Program Pilot untuk Sekolah Anda", 36, WHITE, True, True)
para(t, "Mulai dari satu gerbang, satu kelas, atau satu angkatan — tanpa investasi hardware mahal.", 16, RGBColor(0x94, 0xA3, 0xB8))
tb = box(s, 0.8, 3.5, 11.7, 2.2)
for x in [
    "✓  Setup dibantu penuh: impor data, enroll wajah, registrasi kartu",
    "✓  Tablet/HP bekas pun bisa jadi kiosk — hemat jutaan rupiah",
    "✓  Billing transparan per jumlah siswa aktif per bulan",
    "✓  Dukungan dwibahasa & pelatihan admin/guru",
]:
    para(tb, x, 17, WHITE, first=(x.startswith("✓  Setup")))
t = box(s, 0.8, 6.1, 11.7, 1.0)
para(t, "Hubungi kami untuk demo gratis di sekolah Anda", 22, TEAL_LIGHT, True, PP_ALIGN.CENTER, True)
para(t, "RadiusGate — radiusgate.id", 14, RGBColor(0x94, 0xA3, 0xB8), False, PP_ALIGN.CENTER)

prs.save("/app/frontend/presentasi-radiusgate.pptx")
print("SAVED", len(prs.slides.__iter__.__self__._sldIdLst), "slides")
