import base64
import calendar
import io
import os
import re
import uuid
import zipfile
from datetime import datetime, timezone
from zoneinfo import ZoneInfo
import pandas as pd
from fastapi import APIRouter, HTTPException, Depends, UploadFile, File
from fastapi.responses import StreamingResponse, FileResponse, Response
from pydantic import BaseModel, EmailStr
from typing import List
from db import db
from auth import require_roles, hash_password
from faceutil import embed, cos_sim, NoFaceError, MATCH_SIM_THRESHOLD
from starlette.concurrency import run_in_threadpool
from pdfgen import build_report_pdf, build_kiosk_poster_pdf, build_recap_pdf, build_warning_letter_pdf
from storage import put_object, get_object
from notif import normalize_phone, send_whatsapp

router = APIRouter(tags=["admin"])
admin_dep = require_roles("school_admin")


def now_iso():
    return datetime.now(timezone.utc).isoformat()


async def school_today(sid: str) -> str:
    st = await db.settings.find_one({"school_id": sid}, {"_id": 0, "timezone": 1})
    try:
        tz = ZoneInfo((st or {}).get("timezone", "Asia/Jakarta"))
    except Exception:
        tz = ZoneInfo("Asia/Jakarta")
    return datetime.now(timezone.utc).astimezone(tz).date().isoformat()


@router.get("/admin/subject-att/recap-export")
async def admin_subject_recap_export(month: str, user: dict = Depends(admin_dep)):
    import re as _re
    import calendar as _cal
    from openpyxl import Workbook
    from fastapi.responses import StreamingResponse
    if not _re.match(r"^\d{4}-\d{2}$", month or ""):
        raise HTTPException(status_code=422, detail="invalid_month")
    y, m = map(int, month.split("-"))
    ndays = _cal.monthrange(y, m)[1]
    dates = [f"{month}-{d:02d}" for d in range(1, ndays + 1)]
    recs = await db.subject_attendance.find(
        {"school_id": user["school_id"], "date": {"$regex": f"^{month}-"}},
        {"_id": 0, "student_id": 1, "student_name": 1, "nis": 1, "date": 1, "status": 1,
         "subject": 1, "class_name": 1, "teacher_name": 1}).to_list(100000)
    if not recs:
        raise HTTPException(status_code=404, detail="no_data")
    groups = {}
    for r in recs:
        gkey = (r.get("subject") or "-", r.get("class_name") or "-", r.get("teacher_name") or "-")
        g = groups.setdefault(gkey, {})
        g.setdefault(r["student_id"], {"name": r.get("student_name", ""), "nis": r.get("nis", ""), "days": {}})["days"][r["date"]] = r["status"]
    abbr = {"hadir": "H", "sakit": "S", "izin": "I", "alpha": "A"}
    wb = Workbook()
    first = True
    for (subject, cls, tname), per in sorted(groups.items()):
        ws = wb.active if first else wb.create_sheet()
        first = False
        base = _re.sub(r"[\\/*?:\[\]]", "-", f"{subject} - {cls}")[:31]
        title = base
        n = 1
        while title in wb.sheetnames:
            n += 1
            title = f"{base[:27]} ({n})"
        ws.title = title
        ws.append([f"Rekap Absensi Mapel Bulanan: {subject} — {cls} (Guru: {tname})"])
        ws.append(["Bulan", month])
        ws.append([])
        ws.append(["No", "Nama Siswa", "NIS"] + [str(d) for d in range(1, ndays + 1)] + ["H", "S", "I", "A"])
        for i, st in enumerate(sorted(per.values(), key=lambda x: x["name"]), 1):
            counts = {"hadir": 0, "sakit": 0, "izin": 0, "alpha": 0}
            row = [i, st["name"], st["nis"]]
            for d in dates:
                v = st["days"].get(d, "")
                row.append(abbr.get(v, ""))
                if v in counts:
                    counts[v] += 1
            ws.append(row + [counts["hadir"], counts["sakit"], counts["izin"], counts["alpha"]])
        ws.column_dimensions["A"].width = 6
        ws.column_dimensions["B"].width = 30
        ws.column_dimensions["C"].width = 14
    bio = io.BytesIO()
    wb.save(bio)
    bio.seek(0)
    fname = f"rekap-mapel-bulanan-{month}.xlsx"
    return StreamingResponse(bio, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                             headers={"Content-Disposition": f'attachment; filename="{fname}"'})


# ---------- QR Code ----------
@router.get("/admin/qrcodes/{ptype}/{pid}")
async def get_qrcode(ptype: str, pid: str, user: dict = Depends(admin_dep)):
    coll = {"student": db.students, "teacher": db.teachers, "employee": db.employees}.get(ptype)
    if coll is None:
        raise HTTPException(status_code=404, detail="not_found")
    flt = {"id": pid, "school_id": user["school_id"]}
    person = await coll.find_one(flt, {"_id": 0})
    if not person:
        raise HTTPException(status_code=404, detail="not_found")
    tok = person.get("qr_token")
    if not tok:
        tok = uuid.uuid4().hex
        await coll.update_one(flt, {"$set": {"qr_token": tok}})
    return {"qr": f"RG1.{ptype}.{pid}.{tok}", "name": person["name"],
            "nis": person.get("nis") or person.get("nip") or "", "photo": person.get("photo") or ""}


@router.get("/admin/qrcodes-pdf")
async def qrcodes_pdf(class_name: str = "", user: dict = Depends(admin_dep)):
    import qrcode
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.units import mm
    from reportlab.pdfgen import canvas as rl_canvas
    from reportlab.lib.utils import ImageReader
    from fastapi.responses import StreamingResponse

    flt = {"school_id": user["school_id"], "status": {"$ne": "lulus"}}
    if class_name.strip():
        flt["class"] = class_name.strip()
    students = await db.students.find(flt, {"_id": 0, "embedding": 0, "photo": 0}).sort([("class", 1), ("name", 1)]).to_list(2000)
    if not students:
        raise HTTPException(status_code=404, detail="no_data")
    for st in students:
        if not st.get("qr_token"):
            st["qr_token"] = uuid.uuid4().hex
            await db.students.update_one({"id": st["id"]}, {"$set": {"qr_token": st["qr_token"]}})

    school = await db.schools.find_one({"id": user["school_id"]}, {"_id": 0, "name": 1})
    school_name = (school or {}).get("name", "")

    buf = io.BytesIO()
    c = rl_canvas.Canvas(buf, pagesize=A4)
    W, H = A4
    cols, rows = 2, 3  # 6 kartu per halaman, QR besar agar mudah discan
    cw, ch = W / cols, H / rows
    for i, st in enumerate(students):
        idx = i % (cols * rows)
        if i and idx == 0:
            c.showPage()
        col, row = idx % cols, idx // cols
        x, y = col * cw, H - (row + 1) * ch
        c.setStrokeColorRGB(0.7, 0.7, 0.7)
        c.rect(x + 4 * mm, y + 4 * mm, cw - 8 * mm, ch - 8 * mm)
        c.setFillColorRGB(0.06, 0.42, 0.38)
        c.setFont("Helvetica-Bold", 8)
        c.drawCentredString(x + cw / 2, y + ch - 9 * mm, school_name[:40])
        qr_img = qrcode.make(f"RG1.student.{st['id']}.{st['qr_token']}")
        ib = io.BytesIO()
        qr_img.save(ib, format="PNG")
        ib.seek(0)
        qs = 62 * mm
        c.drawImage(ImageReader(ib), x + (cw - qs) / 2, y + ch - qs - 13 * mm, qs, qs)
        c.setFillColorRGB(0, 0, 0)
        c.setFont("Helvetica-Bold", 11)
        c.drawCentredString(x + cw / 2, y + 11 * mm, st["name"][:30])
        c.setFont("Helvetica", 8)
        c.drawCentredString(x + cw / 2, y + 6.5 * mm, f"NIS: {st.get('nis') or '-'} · {st.get('class') or '-'}")
    c.save()
    buf.seek(0)
    fname = f"kartu-qr-{(class_name.strip() or 'semua').replace(' ', '-')}.pdf"
    return StreamingResponse(buf, media_type="application/pdf",
                             headers={"Content-Disposition": f'attachment; filename="{fname}"'})


# ---------- Dashboard ----------
@router.get("/admin/stats")
async def stats(user: dict = Depends(admin_dep)):
    sid = user["school_id"]
    today = await school_today(sid)
    today_att = await db.attendance.find({"school_id": sid, "date": today, "type": "in"}, {"_id": 0}).to_list(5000)
    students_pc = await db.students.find({"school_id": sid, "status": {"$ne": "lulus"}},
                                         {"_id": 0, "class": 1, "parent_phone": 1, "parent_email": 1,
                                          "parent_name": 1, "address": 1}).to_list(5000)
    per_class = {}
    tot_phone = tot_complete = 0
    for s in students_pc:
        cls = s.get("class") or "-"
        d = per_class.setdefault(cls, {"class": cls, "total": 0, "phone": 0, "complete": 0})
        has_phone = bool(s.get("parent_phone"))
        complete = bool(s.get("parent_phone") and s.get("parent_email") and s.get("parent_name") and s.get("address"))
        d["total"] += 1
        d["phone"] += 1 if has_phone else 0
        d["complete"] += 1 if complete else 0
        tot_phone += 1 if has_phone else 0
        tot_complete += 1 if complete else 0
    parent_data = {
        "total": len(students_pc), "with_phone": tot_phone, "with_complete": tot_complete,
        "per_class": sorted(per_class.values(), key=lambda x: x["class"]),
    }
    return {
        "present_today": len({a["teacher_id"] for a in today_att if a.get("person_type", "teacher") == "teacher"}),
        "students_present": len({a.get("student_id") for a in today_att
                                 if a.get("person_type") == "student" and a.get("att_status", "present") == "present"}),
        "late_today": len([a for a in today_att if a.get("status") == "late"]),
        "pending_leaves": await db.leaves.count_documents({"school_id": sid, "status": "pending"}),
        "total_teachers": await db.teachers.count_documents({"school_id": sid, "active": True}),
        "total_students": await db.students.count_documents({"school_id": sid, "status": {"$ne": "lulus"}}),
        "employees_present": len({a.get("employee_id") for a in today_att if a.get("person_type") == "employee"}),
        "total_employees": await db.employees.count_documents({"school_id": sid, "active": True}),
        "pending_overtime": await db.overtime_requests.count_documents({"school_id": sid, "status": "pending"}),
        "parent_data": parent_data,
    }


@router.get("/admin/today")
async def today_list(date: str | None = None, user: dict = Depends(admin_dep)):
    if date and not re.match(r"^\d{4}-\d{2}-\d{2}$", date):
        raise HTTPException(status_code=422, detail="Format tanggal tidak valid")
    d = date or await school_today(user["school_id"])
    return await db.attendance.find(
        {"school_id": user["school_id"], "date": d}, {"_id": 0, "photo": 0}
    ).sort("ts_server", -1).to_list(500)


@router.delete("/admin/attendance/{aid}")
async def delete_attendance(aid: str, user: dict = Depends(admin_dep)):
    res = await db.attendance.delete_one({"id": aid, "school_id": user["school_id"]})
    if res.deleted_count == 0:
        raise HTTPException(status_code=404, detail="Data absensi tidak ditemukan")
    return {"ok": True}


class AttStatusIn(BaseModel):
    att_status: str  # present|sakit|izin|alpa


@router.patch("/admin/attendance/{aid}/status")
async def admin_set_att_status(aid: str, body: AttStatusIn, user: dict = Depends(admin_dep)):
    if body.att_status not in ("present", "sakit", "izin", "alpa"):
        raise HTTPException(status_code=422, detail="invalid_status")
    upd = {"att_status": body.att_status, "corrected_by": user["email"]}
    if body.att_status in ("sakit", "izin", "alpa"):
        upd.update({"late_minutes": 0, "overtime_minutes": 0, "status": "ok"})
    r = await db.attendance.update_one({"id": aid, "school_id": user["school_id"]}, {"$set": upd})
    if not r.matched_count:
        raise HTTPException(status_code=404, detail="not_found")
    return {"ok": True}


class AttMarkIn(BaseModel):
    student_id: str
    date: str
    att_status: str


@router.post("/admin/attendance/mark")
async def admin_mark_student(body: AttMarkIn, user: dict = Depends(admin_dep)):
    from routes_kiosk import _record  # impor lokal: hindari circular import
    if body.att_status not in ("present", "sakit", "izin", "alpa"):
        raise HTTPException(status_code=422, detail="invalid_status")
    st = await db.students.find_one({"id": body.student_id, "school_id": user["school_id"]},
                                    {"_id": 0, "embedding": 0})
    if not st:
        raise HTTPException(status_code=404, detail="student_not_found")
    existing = await db.attendance.find_one(
        {"school_id": user["school_id"], "student_id": st["id"], "date": body.date, "type": "in"})
    if existing:
        await db.attendance.update_one({"id": existing["id"]},
                                       {"$set": {"att_status": body.att_status, "corrected_by": user["email"]}})
        return {"ok": True, "updated": True}
    await _record({"id": user["school_id"]}, st["id"], st["name"], "in", f"{body.date}T07:00:00",
                  0, 0, "", uuid.uuid4().hex, offline=False, manual=True,
                  extra={"person_type": "student", "class": st.get("class", ""),
                         "att_status": body.att_status, "note": "Ditandai admin",
                         "recorded_by_name": user["name"]})
    return {"ok": True, "created": True}


@router.get("/admin/today/absent")
async def admin_today_absent(date: str, user: dict = Depends(admin_dep)):
    students = await db.students.find(
        {"school_id": user["school_id"], "status": {"$ne": "lulus"}},
        {"_id": 0, "id": 1, "name": 1, "class": 1}).sort("name", 1).to_list(5000)
    have = {r["student_id"] for r in await db.attendance.find(
        {"school_id": user["school_id"], "person_type": "student", "date": date, "type": "in"},
        {"_id": 0, "student_id": 1}).to_list(10000)}
    return [s for s in students if s["id"] not in have]


# ---------- Teachers ----------
def _split_csv(s) -> list:
    return [p.strip() for p in (s or "").split(",") if p.strip()]


async def _derived_list(sid: str, kind: str) -> list:
    if kind == "major":
        return []
    if kind == "department":
        raw = await db.employees.distinct("department", {"school_id": sid})
        return sorted(d for d in raw if d)
    if kind == "class":
        raw = await db.students.distinct("class", {"school_id": sid, "status": {"$ne": "lulus"}})
        return sorted(c for c in raw if c)
    raw = await db.teachers.distinct("subject", {"school_id": sid})
    return sorted({p for s in raw for p in _split_csv(s)})


async def _effective_list(sid: str, kind: str) -> tuple[list, str]:
    key = {"class": "class_list", "subject": "subject_list", "major": "major_list", "department": "department_list"}[kind]
    st = await db.settings.find_one({"school_id": sid}, {"_id": 0, key: 1}) or {}
    if st.get(key) is not None:
        return list(st[key]), key
    return await _derived_list(sid, kind), key


@router.get("/admin/meta/options")
async def meta_options(user: dict = Depends(admin_dep)):
    """Opsi checkbox form guru/siswa: daftar master jika diset, selain itu diturunkan dari data."""
    classes, _ = await _effective_list(user["school_id"], "class")
    subjects, _ = await _effective_list(user["school_id"], "subject")
    majors, _ = await _effective_list(user["school_id"], "major")
    departments, _ = await _effective_list(user["school_id"], "department")
    # mapel yang punya data absensi (mis. mapel kustom yang belum masuk master list) harus selalu muncul di filter
    att_subjects = await db.subject_attendance.distinct("subject", {"school_id": user["school_id"]})
    subjects = sorted(set(subjects) | {s for s in att_subjects if s})
    # hal yang sama untuk kelas: kelas yang punya data absensi harus selalu muncul
    att_classes = set(await db.subject_attendance.distinct("class_name", {"school_id": user["school_id"]}))
    att_classes |= set(await db.attendance.distinct("class", {"school_id": user["school_id"]}))
    classes = sorted(set(classes) | {c for c in att_classes if c})
    return {"classes": classes, "subjects": subjects, "majors": majors, "departments": departments}


class MetaRenameIn(BaseModel):
    kind: str  # class | subject
    from_value: str
    to_value: str


class MetaDeleteIn(BaseModel):
    kind: str
    value: str


@router.post("/admin/meta/rename")
async def meta_rename(body: MetaRenameIn, user: dict = Depends(admin_dep)):
    """Ganti nama kelas/mapel/departemen, otomatis diterapkan ke siswa, absensi, guru, karyawan."""
    if body.kind not in ("class", "subject", "major", "department"):
        raise HTTPException(status_code=400, detail="kind tidak valid")
    fv, tv = body.from_value.strip(), body.to_value.strip()
    if not fv or not tv or fv == tv:
        raise HTTPException(status_code=400, detail="Nama tidak valid")
    sid = user["school_id"]
    items, key = await _effective_list(sid, body.kind)
    if fv not in items:
        raise HTTPException(status_code=404, detail="Data tidak ditemukan")
    if tv in items:
        raise HTTPException(status_code=400, detail="Nama baru sudah ada")
    items = sorted(tv if x == fv else x for x in items)
    await db.settings.update_one({"school_id": sid}, {"$set": {key: items}}, upsert=True)
    if body.kind == "class":
        await db.students.update_many({"school_id": sid, "class": fv}, {"$set": {"class": tv}})
        await db.attendance.update_many({"school_id": sid, "class": fv}, {"$set": {"class": tv}})
        teachers = await db.teachers.find({"school_id": sid}, {"_id": 0, "id": 1, "classes": 1}).to_list(2000)
        for tch in teachers:
            parts = _split_csv(tch.get("classes"))
            if fv in parts:
                await db.teachers.update_one({"id": tch["id"]},
                                             {"$set": {"classes": ", ".join(tv if p == fv else p for p in parts)}})
    elif body.kind == "department":
        await db.employees.update_many({"school_id": sid, "department": fv}, {"$set": {"department": tv}})
    else:
        teachers = await db.teachers.find({"school_id": sid}, {"_id": 0, "id": 1, "subject": 1}).to_list(2000)
        for tch in teachers:
            parts = _split_csv(tch.get("subject"))
            if fv in parts:
                await db.teachers.update_one({"id": tch["id"]},
                                             {"$set": {"subject": ", ".join(tv if p == fv else p for p in parts)}})
    return {"ok": True, "items": items}


@router.post("/admin/meta/delete")
async def meta_delete(body: MetaDeleteIn, user: dict = Depends(admin_dep)):
    """Hapus item master; ditolak jika masih dipakai data siswa/guru/karyawan."""
    if body.kind not in ("class", "subject", "major", "department"):
        raise HTTPException(status_code=400, detail="kind tidak valid")
    v = body.value.strip()
    sid = user["school_id"]
    if body.kind == "class":
        used = await db.students.count_documents({"school_id": sid, "class": v, "status": {"$ne": "lulus"}})
        if used:
            raise HTTPException(status_code=400, detail=f"class_in_use:{used}")
    elif body.kind == "department":
        used = await db.employees.count_documents({"school_id": sid, "department": v})
        if used:
            raise HTTPException(status_code=400, detail=f"department_in_use:{used}")
    elif body.kind == "subject":
        teachers = await db.teachers.find({"school_id": sid}, {"_id": 0, "subject": 1}).to_list(2000)
        if any(v in _split_csv(tch.get("subject")) for tch in teachers):
            raise HTTPException(status_code=400, detail="subject_in_use")
    items, key = await _effective_list(sid, body.kind)
    items = [x for x in items if x != v]
    await db.settings.update_one({"school_id": sid}, {"$set": {key: items}}, upsert=True)
    return {"ok": True, "items": items}


class TeacherIn(BaseModel):
    name: str
    email: EmailStr
    password: str
    nip: str = ""
    subject: str = ""
    classes: str = ""
    gender: str = ""
    card_uid: str = ""


class TeacherPatch(BaseModel):
    name: str | None = None
    nip: str | None = None
    subject: str | None = None
    active: bool | None = None
    classes: str | None = None
    password: str | None = None
    gender: str | None = None
    card_uid: str | None = None


@router.get("/admin/teachers")
async def list_teachers(user: dict = Depends(admin_dep)):
    teachers = await db.teachers.find({"school_id": user["school_id"]}, {"_id": 0, "embedding": 0, "photo": 0}).to_list(1000)
    for t in teachers:
        u = await db.users.find_one({"id": t["user_id"]}, {"_id": 0, "email": 1})
        t["email"] = u["email"] if u else ""
        enrolled = await db.teachers.find_one({"id": t["id"]}, {"embedding": 1})
        t["enrolled"] = bool(enrolled and enrolled.get("embedding"))
    return teachers


@router.post("/admin/teachers")
async def create_teacher(body: TeacherIn, user: dict = Depends(admin_dep)):
    if await db.users.find_one({"email": body.email.lower()}):
        raise HTTPException(status_code=400, detail="Email sudah dipakai")
    uid = str(uuid.uuid4())
    await db.users.insert_one({
        "id": uid, "email": body.email.lower(), "name": body.name, "role": "teacher",
        "password_hash": hash_password(body.password), "school_id": user["school_id"], "created_at": now_iso(),
    })
    teacher = {
        "id": str(uuid.uuid4()), "school_id": user["school_id"], "user_id": uid,
        "name": body.name, "nip": body.nip, "subject": body.subject, "classes": body.classes,
        "gender": _norm_gender(body.gender), "card_uid": body.card_uid.strip(),
        "embedding": None, "photo": None, "active": True, "created_at": now_iso(),
    }
    await db.teachers.insert_one(teacher)
    teacher.pop("_id", None)
    teacher.pop("embedding", None)
    teacher.pop("photo", None)
    teacher["email"] = body.email.lower()
    teacher["enrolled"] = False
    return teacher


@router.patch("/admin/teachers/{tid}")
async def update_teacher(tid: str, body: TeacherPatch, user: dict = Depends(admin_dep)):
    upd = {k: v for k, v in body.model_dump().items() if v is not None}
    if "gender" in upd:
        upd["gender"] = _norm_gender(upd["gender"])
    new_pw = upd.pop("password", None)
    if new_pw is not None and len(new_pw) < 6:
        raise HTTPException(status_code=422, detail="password_too_short")
    if not upd and not new_pw:
        raise HTTPException(status_code=400, detail="Tidak ada perubahan")
    if upd:
        await db.teachers.update_one({"id": tid, "school_id": user["school_id"]}, {"$set": upd})
    t = await db.teachers.find_one({"id": tid, "school_id": user["school_id"]}, {"_id": 0, "user_id": 1})
    if not t:
        raise HTTPException(status_code=404, detail="Guru tidak ditemukan")
    if "name" in upd:
        await db.users.update_one({"id": t["user_id"]}, {"$set": {"name": upd["name"]}})
    if new_pw:
        await db.users.update_one({"id": t["user_id"]}, {"$set": {"password_hash": hash_password(new_pw)}})
    return {"ok": True}


@router.delete("/admin/teachers/{tid}")
async def delete_teacher(tid: str, user: dict = Depends(admin_dep)):
    t = await db.teachers.find_one({"id": tid, "school_id": user["school_id"]})
    if not t:
        raise HTTPException(status_code=404, detail="Guru tidak ditemukan")
    await db.teachers.delete_one({"id": tid})
    await db.users.delete_one({"id": t["user_id"]})
    return {"ok": True}


class EnrollIn(BaseModel):
    photo: str


async def _face_dup_name(school_id: str, emb, exclude_id: str):
    """Nama orang lain yang wajahnya sudah terdaftar, atau None jika aman."""
    others = await db.teachers.find(
        {"school_id": school_id, "id": {"$ne": exclude_id}, "embedding": {"$type": "array"}},
        {"_id": 0, "name": 1, "embedding": 1}).to_list(2000)
    others += await db.students.find(
        {"school_id": school_id, "id": {"$ne": exclude_id}, "embedding": {"$type": "array"}},
        {"_id": 0, "name": 1, "embedding": 1}).to_list(5000)
    others += await db.employees.find(
        {"school_id": school_id, "id": {"$ne": exclude_id}, "embedding": {"$type": "array"}},
        {"_id": 0, "name": 1, "embedding": 1}).to_list(2000)
    for o in others:
        if cos_sim(emb, o["embedding"]) >= MATCH_SIM_THRESHOLD:
            return o["name"]
    return None


async def _face_dup_check(school_id: str, emb, exclude_id: str):
    dup = await _face_dup_name(school_id, emb, exclude_id)
    if dup:
        raise HTTPException(status_code=409, detail=f"face_already_enrolled:{dup}")


@router.post("/admin/teachers/{tid}/enroll")
async def enroll_face(tid: str, body: EnrollIn, user: dict = Depends(admin_dep)):
    t = await db.teachers.find_one({"id": tid, "school_id": user["school_id"]})
    if not t:
        raise HTTPException(status_code=404, detail="Guru tidak ditemukan")
    try:
        emb = await run_in_threadpool(embed, body.photo)
    except NoFaceError as e:
        raise HTTPException(status_code=422, detail=str(e))
    except Exception:
        raise HTTPException(status_code=400, detail="Foto tidak valid")
    await _face_dup_check(user["school_id"], emb, tid)
    await db.teachers.update_one({"id": tid}, {"$set": {"embedding": emb, "embedding_model": "buffalo_s", "photo": body.photo, "enrolled_at": now_iso()}})
    return {"ok": True, "enrolled": True}


# ---------- Settings & Locations ----------
class SettingsIn(BaseModel):
    work_start: str | None = None
    work_end: str | None = None
    late_tolerance_min: int | None = None
    early_checkin_min: int | None = None
    timezone: str | None = None
    require_checkin: bool | None = None
    greeting_in: str | None = None
    greeting_out: str | None = None
    class_list: list[str] | None = None
    subject_list: list[str] | None = None
    major_list: list[str] | None = None
    school_type: str | None = None
    department_list: list[str] | None = None
    overtime_rate: int | None = None
    saver_notes: list[dict] | None = None
    saver_photos: list[str] | None = None
    saver_enabled: bool | None = None
    saver_photos_enabled: bool | None = None
    student_dismissal: dict | None = None
    org_type: str | None = None
    sp_thresholds: dict | None = None


@router.get("/admin/settings")
async def get_settings(user: dict = Depends(admin_dep)):
    st = await db.settings.find_one({"school_id": user["school_id"]}, {"_id": 0})
    locs = await db.locations.find({"school_id": user["school_id"]}, {"_id": 0}).to_list(100)
    school = await db.schools.find_one({"id": user["school_id"]}, {"_id": 0, "kiosk_token": 1, "name": 1, "org_type": 1})
    return {"settings": st, "locations": locs, "school": school}


@router.put("/admin/settings")
async def put_settings(body: SettingsIn, user: dict = Depends(admin_dep)):
    upd = {k: v for k, v in body.model_dump().items() if v is not None}
    org = upd.pop("org_type", None)
    if org is not None:
        if org not in ("school", "company"):
            raise HTTPException(status_code=422, detail="org_type tidak valid")
        await db.schools.update_one({"id": user["school_id"]}, {"$set": {"org_type": org}})
    if not upd and org is None:
        raise HTTPException(status_code=400, detail="Tidak ada perubahan")
    if upd:
        await db.settings.update_one({"school_id": user["school_id"]}, {"$set": upd}, upsert=True)
    return {"ok": True}


# ---------- Shift Kerja ----------
def _valid_hhmm(v: str):
    try:
        hh, mm = map(int, v.split(":"))
        if not (0 <= hh <= 23 and 0 <= mm <= 59):
            raise ValueError
    except Exception:
        raise HTTPException(status_code=422, detail="Format jam harus HH:MM")


class ShiftIn(BaseModel):
    name: str
    start: str
    end: str


@router.get("/admin/shifts")
async def list_shifts(user: dict = Depends(admin_dep)):
    return await db.shifts.find({"school_id": user["school_id"]}, {"_id": 0}).to_list(100)


@router.post("/admin/shifts")
async def create_shift(body: ShiftIn, user: dict = Depends(admin_dep)):
    _valid_hhmm(body.start)
    _valid_hhmm(body.end)
    doc = {"id": str(uuid.uuid4()), "school_id": user["school_id"], "name": body.name.strip(),
           "start": body.start, "end": body.end, "created_at": now_iso()}
    await db.shifts.insert_one(doc)
    doc.pop("_id", None)
    return doc


@router.patch("/admin/shifts/{shid}")
async def update_shift(shid: str, body: ShiftIn, user: dict = Depends(admin_dep)):
    _valid_hhmm(body.start)
    _valid_hhmm(body.end)
    r = await db.shifts.update_one({"id": shid, "school_id": user["school_id"]},
                                   {"$set": {"name": body.name.strip(), "start": body.start, "end": body.end}})
    if not r.matched_count:
        raise HTTPException(status_code=404, detail="Shift tidak ditemukan")
    return {"ok": True}


@router.delete("/admin/shifts/{shid}")
async def delete_shift(shid: str, user: dict = Depends(admin_dep)):
    await db.shifts.delete_one({"id": shid, "school_id": user["school_id"]})
    await db.employees.update_many({"school_id": user["school_id"], "shift_id": shid}, {"$unset": {"shift_id": ""}})
    return {"ok": True}


# ---------- Surat Peringatan (SP1/SP2/SP3) ----------
def _sp_thresholds(settings: dict | None) -> dict:
    raw = (settings or {}).get("sp_thresholds") or {}
    return {"SP1": int(raw.get("SP1", 3)), "SP2": int(raw.get("SP2", 6)), "SP3": int(raw.get("SP3", 10))}


@router.get("/admin/warnings/candidates")
async def warning_candidates(month: str, user: dict = Depends(admin_dep)):
    sid = user["school_id"]
    settings = await db.settings.find_one({"school_id": sid}, {"_id": 0, "sp_thresholds": 1})
    th = _sp_thresholds(settings)
    rows = await db.attendance.find(
        {"school_id": sid, "date": {"$regex": f"^{month}"}, "status": "late", "type": "in"},
        {"_id": 0, "person_type": 1, "teacher_id": 1, "employee_id": 1, "teacher_name": 1}).to_list(20000)
    counts, meta = {}, {}
    for r in rows:
        pt = r.get("person_type", "teacher")
        pid = r.get("employee_id") if pt == "employee" else r.get("teacher_id")
        if not pid:
            continue
        counts[pid] = counts.get(pid, 0) + 1
        meta[pid] = (r.get("teacher_name") or "", pt)
    issued = await db.warning_letters.find({"school_id": sid, "month": month}, {"_id": 0, "person_id": 1, "level": 1}).to_list(500)
    issued_by = {}
    for wl in issued:
        issued_by.setdefault(wl["person_id"], []).append(wl["level"])
    out = []
    for pid, cnt in counts.items():
        if cnt < th["SP1"]:
            continue
        name, pt = meta[pid]
        levels = issued_by.get(pid, [])
        if "SP3" in levels:
            suggested = None
        elif "SP2" in levels:
            suggested = "SP3" if cnt >= th["SP3"] else None
        elif "SP1" in levels:
            suggested = "SP2" if cnt >= th["SP2"] else None
        else:
            suggested = "SP1"
        out.append({"person_id": pid, "name": name, "person_type": pt, "late_count": cnt,
                    "issued_levels": levels, "suggested": suggested})
    out.sort(key=lambda x: -x["late_count"])
    return {"thresholds": th, "candidates": out}


class WarningIssueIn(BaseModel):
    person_id: str
    person_type: str = "employee"
    month: str
    level: str


@router.post("/admin/warnings/issue")
async def issue_warning(body: WarningIssueIn, user: dict = Depends(admin_dep)):
    if body.level not in ("SP1", "SP2", "SP3"):
        raise HTTPException(status_code=422, detail="Level SP tidak valid")
    if body.person_type not in ("employee", "teacher"):
        raise HTTPException(status_code=422, detail="person_type tidak valid")
    sid = user["school_id"]
    coll = db.employees if body.person_type == "employee" else db.teachers
    person = await coll.find_one({"id": body.person_id, "school_id": sid},
                                 {"_id": 0, "name": 1, "nip": 1, "department": 1, "position": 1})
    if not person:
        raise HTTPException(status_code=404, detail="Karyawan tidak ditemukan")
    if await db.warning_letters.find_one({"school_id": sid, "person_id": body.person_id, "month": body.month, "level": body.level}):
        raise HTTPException(status_code=409, detail="SP level ini sudah diterbitkan untuk periode tersebut")
    id_field = "employee_id" if body.person_type == "employee" else "teacher_id"
    late_count = await db.attendance.count_documents({
        "school_id": sid, "date": {"$regex": f"^{body.month}"}, "status": "late", "type": "in", id_field: body.person_id})
    doc = {"id": str(uuid.uuid4()), "school_id": sid, "person_id": body.person_id,
           "person_type": body.person_type, "name": person["name"],
           "nip": person.get("nip", ""), "department": person.get("department", ""), "position": person.get("position", ""),
           "month": body.month, "level": body.level, "late_count": late_count,
           "issued_by": user.get("name", ""), "created_at": now_iso()}
    await db.warning_letters.insert_one(doc)
    doc.pop("_id", None)
    return doc


@router.get("/admin/warnings")
async def list_warnings(month: str | None = None, user: dict = Depends(admin_dep)):
    q = {"school_id": user["school_id"]}
    if month:
        q["month"] = month
    return await db.warning_letters.find(q, {"_id": 0}).sort("created_at", -1).to_list(500)


@router.get("/admin/warnings/{wid}/pdf")
async def warning_pdf(wid: str, user: dict = Depends(admin_dep)):
    wl = await db.warning_letters.find_one({"id": wid, "school_id": user["school_id"]}, {"_id": 0})
    if not wl:
        raise HTTPException(status_code=404, detail="Surat tidak ditemukan")
    school = await db.schools.find_one({"id": user["school_id"]}, {"_id": 0, "name": 1, "address": 1})
    path = build_warning_letter_pdf(wl, school or {})
    return FileResponse(path, media_type="application/pdf", filename=f"{wl['level']}-{wl['name']}-{wl['month']}.pdf")


# ---------- Foto slide screensaver kiosk (object storage) ----------
@router.post("/admin/saver-photos")
async def upload_saver_photo(file: UploadFile = File(...), user: dict = Depends(admin_dep)):
    if file.content_type not in ("image/jpeg", "image/png", "image/webp"):
        raise HTTPException(status_code=400, detail="Format harus JPG/PNG/WEBP")
    data = await file.read()
    if len(data) > 2 * 1024 * 1024:
        raise HTTPException(status_code=400, detail="Ukuran maksimal 2MB")
    ext = file.filename.rsplit(".", 1)[-1].lower() if "." in file.filename else "jpg"
    path = f"radiusgate/saver/{user['school_id']}/{uuid.uuid4()}.{ext}"
    res = await run_in_threadpool(put_object, path, data, file.content_type)
    await db.settings.update_one({"school_id": user["school_id"]}, {"$push": {"saver_photos": res["path"]}}, upsert=True)
    return {"path": res["path"]}


@router.delete("/admin/saver-photos")
async def delete_saver_photo(path: str, user: dict = Depends(admin_dep)):
    await db.settings.update_one({"school_id": user["school_id"]}, {"$pull": {"saver_photos": path}})
    return {"ok": True}


@router.get("/admin/saver-photos/file/{path:path}")
async def saver_photo_file(path: str):
    data, ct = await run_in_threadpool(get_object, path)
    return Response(content=data, media_type=ct)


@router.get("/admin/kiosk-poster")
async def kiosk_poster(user: dict = Depends(admin_dep)):
    school = await db.schools.find_one({"id": user["school_id"]}, {"_id": 0})
    if not school or not school.get("kiosk_token"):
        raise HTTPException(status_code=404, detail="Sekolah tidak ditemukan")
    pair_url = f"{os.environ.get('FRONTEND_URL', '')}/kiosk?pair={school['kiosk_token']}"
    path = await run_in_threadpool(build_kiosk_poster_pdf, school, pair_url)
    return FileResponse(path, media_type="application/pdf", filename=f"poster-kiosk-{school['kiosk_token']}.pdf")


class LocationIn(BaseModel):
    name: str
    lat: float
    lng: float
    radius_m: int = 50


@router.post("/admin/locations")
async def add_location(body: LocationIn, user: dict = Depends(admin_dep)):
    loc = {"id": str(uuid.uuid4()), "school_id": user["school_id"], **body.model_dump()}
    await db.locations.insert_one(loc)
    loc.pop("_id", None)
    return loc


@router.patch("/admin/locations/{lid}")
async def update_location(lid: str, body: LocationIn, user: dict = Depends(admin_dep)):
    await db.locations.update_one({"id": lid, "school_id": user["school_id"]}, {"$set": body.model_dump()})
    return {"ok": True}


@router.delete("/admin/locations/{lid}")
async def delete_location(lid: str, user: dict = Depends(admin_dep)):
    await db.locations.delete_one({"id": lid, "school_id": user["school_id"]})
    return {"ok": True}


# ---------- Students ----------
class StudentIn(BaseModel):
    name: str
    nis: str = ""
    nisn: str = ""
    gender: str = ""
    class_name: str = ""
    parent_phone: str = ""
    parent_name: str = ""
    parent_email: str = ""
    address: str = ""
    card_uid: str = ""


def _norm_gender(v: str) -> str:
    v = (v or "").strip().lower()
    if v in ("l", "lk", "laki", "laki-laki", "male", "m"):
        return "L"
    if v in ("p", "pr", "perempuan", "female", "f", "w", "wanita"):
        return "P"
    return ""


@router.get("/admin/students")
async def list_students(user: dict = Depends(admin_dep)):
    students = await db.students.find({"school_id": user["school_id"]}, {"_id": 0, "embedding": 0, "photo": 0}).to_list(5000)
    enrolled_ids = {s["id"] for s in await db.students.find(
        {"school_id": user["school_id"], "embedding": {"$ne": None}}, {"_id": 0, "id": 1}).to_list(5000)}
    parent_sids = {u["student_id"] for u in await db.users.find(
        {"school_id": user["school_id"], "role": "parent"}, {"_id": 0, "student_id": 1}).to_list(5000)}
    for s in students:
        s["enrolled"] = s["id"] in enrolled_ids
        s["has_parent_account"] = s["id"] in parent_sids
    return students


@router.post("/admin/students")
async def add_student(body: StudentIn, user: dict = Depends(admin_dep)):
    st = {"id": str(uuid.uuid4()), "school_id": user["school_id"],
          "name": body.name, "nis": body.nis, "nisn": body.nisn,
          "gender": _norm_gender(body.gender), "class": body.class_name, "status": "aktif",
          "parent_phone": normalize_phone(body.parent_phone),
          "parent_name": body.parent_name.strip(), "parent_email": body.parent_email.strip().lower(),
          "address": body.address.strip(), "card_uid": body.card_uid.strip()}
    await db.students.insert_one(st)
    st.pop("_id", None)
    st["parent_account"] = await _ensure_parent_account(user["school_id"], st)
    return st


@router.post("/admin/students/{stid}/enroll")
async def enroll_student_face(stid: str, body: EnrollIn, user: dict = Depends(admin_dep)):
    st = await db.students.find_one({"id": stid, "school_id": user["school_id"]})
    if not st:
        raise HTTPException(status_code=404, detail="Siswa tidak ditemukan")
    try:
        emb = await run_in_threadpool(embed, body.photo)
    except NoFaceError as e:
        raise HTTPException(status_code=422, detail=str(e))
    except Exception:
        raise HTTPException(status_code=400, detail="Foto tidak valid")
    await _face_dup_check(user["school_id"], emb, stid)
    await db.students.update_one({"id": stid}, {"$set": {"embedding": emb, "embedding_model": "buffalo_s", "photo": body.photo, "enrolled_at": now_iso()}})
    return {"ok": True, "enrolled": True}


@router.post("/admin/students/enroll-zip")
async def enroll_students_zip(file: UploadFile = File(...), mapping: UploadFile = File(None), user: dict = Depends(admin_dep)):
    """Enroll wajah massal dari ZIP foto. Nama file = NIS, atau dicocokkan via mapping Excel/CSV (kolom NIS + nama file)."""
    sid = user["school_id"]
    try:
        zf = zipfile.ZipFile(io.BytesIO(await file.read()))
    except Exception:
        raise HTTPException(status_code=400, detail="File bukan ZIP yang valid")
    fmap = {}
    if mapping and mapping.filename:
        mc = await mapping.read()
        try:
            mdf = pd.read_excel(io.BytesIO(mc), dtype=str) if mapping.filename.lower().endswith((".xlsx", ".xls")) else pd.read_csv(io.BytesIO(mc), dtype=str)
        except Exception:
            raise HTTPException(status_code=400, detail="File mapping tidak bisa dibaca")
        mdf.columns = [str(c).strip().lower() for c in mdf.columns]
        nis_col = next((c for c in mdf.columns if "nis" in c and "nisn" not in c), None)
        file_col = next((c for c in mdf.columns if any(k in c for k in ("file", "foto", "photo"))), None)
        if not nis_col or not file_col:
            raise HTTPException(status_code=400, detail="Mapping wajib punya kolom NIS dan kolom nama file/foto")
        for _, r in mdf.iterrows():
            fn, nis = str(r[file_col]).strip().lower(), str(r[nis_col]).strip()
            if fn and nis and fn != "nan" and nis.lower() != "nan":
                fmap[fn] = nis
                fmap[fn.rsplit(".", 1)[0]] = nis
    students = {s["nis"]: s for s in await db.students.find(
        {"school_id": sid, "nis": {"$nin": [None, ""]}}, {"_id": 0, "embedding": 0, "photo": 0}).to_list(10000)}
    entries = [i for i in zf.infolist() if not i.is_dir() and not i.filename.startswith("__MACOSX")
               and i.filename.lower().rsplit(".", 1)[-1] in ("jpg", "jpeg", "png")][:2000]
    results = []
    for info in entries:
        base = info.filename.split("/")[-1]
        stem = base.rsplit(".", 1)[0]
        nis = fmap.get(base.lower()) or fmap.get(stem.lower()) or stem.strip()
        st = students.get(nis)
        if not st:
            results.append({"file": base, "nis": nis, "name": "", "status": "gagal", "reason": "NIS tidak ditemukan di data siswa"})
            continue
        try:
            b64 = "data:image/jpeg;base64," + base64.b64encode(zf.read(info)).decode()
            emb = await run_in_threadpool(embed, b64)
        except NoFaceError as e:
            results.append({"file": base, "nis": nis, "name": st["name"], "status": "gagal",
                            "reason": "Wajah tidak terdeteksi" if "no_face" in str(e) else "Lebih dari 1 wajah di foto"})
            continue
        except Exception:
            results.append({"file": base, "nis": nis, "name": st["name"], "status": "gagal", "reason": "File gambar rusak/tidak bisa dibaca"})
            continue
        dup = await _face_dup_name(sid, emb, st["id"])
        if dup:
            results.append({"file": base, "nis": nis, "name": st["name"], "status": "gagal", "reason": f"Wajah sudah terdaftar atas nama {dup}"})
            continue
        await db.students.update_one({"id": st["id"]}, {"$set": {"embedding": emb, "embedding_model": "buffalo_s", "enrolled_at": now_iso()}})
        results.append({"file": base, "nis": nis, "name": st["name"], "status": "sukses", "reason": ""})
    ok = len([r for r in results if r["status"] == "sukses"])
    return {"total": len(results), "success": ok, "failed": len(results) - ok, "results": results}


class StudentPatch(BaseModel):
    name: str | None = None
    nis: str | None = None
    nisn: str | None = None
    gender: str | None = None
    class_name: str | None = None
    parent_phone: str | None = None
    parent_name: str | None = None
    parent_email: str | None = None
    address: str | None = None
    card_uid: str | None = None


class BulkDeleteIn(BaseModel):
    ids: list[str]


@router.patch("/admin/students/{stid}")
async def update_student(stid: str, body: StudentPatch, user: dict = Depends(admin_dep)):
    upd = {k: v for k, v in body.model_dump(exclude_unset=True).items()}
    if "class_name" in upd:
        upd["class"] = upd.pop("class_name")
    if "parent_phone" in upd:
        upd["parent_phone"] = normalize_phone(upd["parent_phone"])
    if not upd:
        raise HTTPException(status_code=400, detail="Tidak ada perubahan")
    res = await db.students.update_one({"id": stid, "school_id": user["school_id"]}, {"$set": upd})
    if res.matched_count == 0:
        raise HTTPException(status_code=404, detail="Siswa tidak ditemukan")
    if "parent_phone" in upd:
        sync = await db.users.update_one({"role": "parent", "student_id": stid},
                                         {"$set": {"phone": upd["parent_phone"],
                                                   "email": f"ortu+{upd['parent_phone']}@edugateid.local"}})
        if sync.matched_count:
            return {"ok": True, "parent_account": "exists"}
        st2 = await db.students.find_one({"id": stid}, {"_id": 0, "id": 1, "name": 1, "nis": 1, "parent_phone": 1})
        return {"ok": True, "parent_account": await _ensure_parent_account(user["school_id"], st2)}
    return {"ok": True}


@router.post("/admin/students/bulk-delete")
async def bulk_delete_students(body: BulkDeleteIn, user: dict = Depends(admin_dep)):
    ids = body.ids[:1000]
    res = await db.students.delete_many({"id": {"$in": ids}, "school_id": user["school_id"]})
    await db.users.delete_many({"role": "parent", "student_id": {"$in": ids}})
    return {"deleted": res.deleted_count}


class PromoteIn(BaseModel):
    from_class: str
    to_class: str


class GraduateIn(BaseModel):
    class_name: str


@router.post("/admin/students/promote")
async def promote_students(body: PromoteIn, user: dict = Depends(admin_dep)):
    if not body.from_class.strip() or not body.to_class.strip():
        raise HTTPException(status_code=422, detail="empty_class")
    res = await db.students.update_many(
        {"school_id": user["school_id"], "class": body.from_class.strip(), "status": {"$ne": "lulus"}},
        {"$set": {"class": body.to_class.strip()}})
    return {"updated": res.modified_count}


@router.post("/admin/students/graduate")
async def graduate_students(body: GraduateIn, user: dict = Depends(admin_dep)):
    if not body.class_name.strip():
        raise HTTPException(status_code=422, detail="empty_class")
    res = await db.students.update_many(
        {"school_id": user["school_id"], "class": body.class_name.strip(), "status": {"$ne": "lulus"}},
        {"$set": {"status": "lulus"}})
    return {"updated": res.modified_count}


class YearActionIn(BaseModel):
    promote: list[PromoteIn] = []
    graduate: list[str] = []


@router.post("/admin/students/promote-year")
async def promote_year(body: YearActionIn, user: dict = Depends(admin_dep)):
    """Kenaikan kelas massal 1 tahun ajaran. Kelulusan diproses dulu, lalu kenaikan
    diurutkan menurun (XII->XI->X) agar tidak ada siswa yang naik dua kali."""
    graduated = 0
    for cls in body.graduate[:200]:
        if not cls.strip():
            continue
        res = await db.students.update_many(
            {"school_id": user["school_id"], "class": cls.strip(), "status": {"$ne": "lulus"}},
            {"$set": {"status": "lulus"}})
        graduated += res.modified_count
    promoted = 0
    for m in sorted(body.promote[:200], key=lambda x: x.from_class, reverse=True):
        if not m.from_class.strip() or not m.to_class.strip():
            continue
        res = await db.students.update_many(
            {"school_id": user["school_id"], "class": m.from_class.strip(), "status": {"$ne": "lulus"}},
            {"$set": {"class": m.to_class.strip()}})
        promoted += res.modified_count
    return {"promoted": promoted, "graduated": graduated}


@router.delete("/admin/students/{stid}")
async def delete_student(stid: str, user: dict = Depends(admin_dep)):
    await db.students.delete_one({"id": stid, "school_id": user["school_id"]})
    await db.users.delete_many({"role": "parent", "student_id": stid})
    return {"ok": True}


async def _ensure_parent_account(sid: str, st: dict) -> str | None:
    """Buat akun ortu otomatis saat parent_phone diisi. Return 'created'|'exists'|'phone_used'|None."""
    phone = st.get("parent_phone", "")
    if not phone:
        return None
    if await db.users.find_one({"role": "parent", "student_id": st["id"]}):
        return "exists"
    if await db.users.find_one({"role": "parent", "phone": phone}):
        return "phone_used"
    await db.users.insert_one({
        "id": str(uuid.uuid4()), "email": f"ortu+{phone}@edugateid.local",
        "name": st.get("parent_name") or f"Orang Tua {st['name']}", "role": "parent", "phone": phone,
        "student_id": st["id"], "school_id": sid,
        "password_hash": hash_password(st.get("nis") or phone[-6:]),
        "created_at": now_iso(),
    })
    return "created"


@router.post("/admin/students/{stid}/send-parent-login")
async def send_parent_login(stid: str, user: dict = Depends(admin_dep)):
    """Kirim info login portal ortu via WA. Password di-reset ke NIS anak HANYA jika WA benar-benar
    terkirim (Wablas aktif); jika gateway belum aktif, kembalikan wa_link untuk kirim manual."""
    sid = user["school_id"]
    st = await db.students.find_one({"id": stid, "school_id": sid}, {"_id": 0, "embedding": 0, "photo": 0})
    if not st:
        raise HTTPException(status_code=404, detail="Siswa tidak ditemukan")
    phone = normalize_phone(st.get("parent_phone", ""))
    if not phone:
        raise HTTPException(status_code=422, detail="Siswa belum punya No. HP orang tua")
    acct = await _ensure_parent_account(sid, st)
    if acct == "phone_used":
        raise HTTPException(status_code=400, detail="No. HP sudah dipakai akun ortu lain")
    password = st.get("nis") or phone[-6:]
    school = await db.schools.find_one({"id": sid}, {"_id": 0, "name": 1})
    portal = os.environ.get("FRONTEND_URL", "")
    msg = (f"RadiusGate - {(school or {}).get('name', '')}\n"
           f"Info Login Portal Orang Tua untuk memantau absensi & tagihan Ananda *{st['name']}*:\n\n"
           f"Portal: {portal}\nLogin: {phone}\nPassword: {password}\n\n"
           f"Segera ganti password setelah masuk (menu Ganti Password).")
    res = await send_whatsapp(phone, msg)
    sent = res.get("mode") == "wablas" and not res.get("error")
    if sent:
        await db.users.update_one({"role": "parent", "student_id": stid},
                                  {"$set": {"password_hash": hash_password(password)}})
    return {"ok": True, "sent": sent, "password_reset": sent, "wa_link": res.get("wa_link"), "account": acct or "exists"}


@router.post("/admin/students/create-parent-accounts")
async def create_parent_accounts(user: dict = Depends(admin_dep)):
    """Buat akun orang tua massal dari no. HP di data siswa. Login: no. HP, password awal: NIS anak."""
    sid = user["school_id"]
    students = await db.students.find(
        {"school_id": sid, "status": {"$ne": "lulus"}, "parent_phone": {"$nin": [None, ""]}},
        {"_id": 0, "embedding": 0, "photo": 0}).to_list(5000)
    created, skipped = 0, []
    for s in students:
        phone = normalize_phone(s.get("parent_phone", ""))
        if not phone:
            skipped.append({"name": s["name"], "reason": "hp_kosong"})
            continue
        if await db.users.find_one({"role": "parent", "student_id": s["id"]}):
            skipped.append({"name": s["name"], "reason": "sudah_ada"})
            continue
        if await db.users.find_one({"role": "parent", "phone": phone}):
            skipped.append({"name": s["name"], "reason": "hp_dipakai"})
            continue
        await db.users.insert_one({
            "id": str(uuid.uuid4()), "email": f"ortu+{phone}@edugateid.local",
            "name": s.get("parent_name") or f"Orang Tua {s['name']}", "role": "parent", "phone": phone,
            "student_id": s["id"], "school_id": sid,
            "password_hash": hash_password(s.get("nis") or phone[-6:]),
            "created_at": now_iso(),
        })
        created += 1
    return {"created": created, "skipped": skipped}


@router.post("/admin/students/import/preview")
async def import_preview(file: UploadFile = File(...), user: dict = Depends(admin_dep)):
    content = await file.read()
    try:
        if file.filename.lower().endswith((".xlsx", ".xls")):
            df = pd.read_excel(io.BytesIO(content), dtype=str)
        else:
            df = pd.read_csv(io.BytesIO(content), dtype=str)
    except Exception:
        raise HTTPException(status_code=400, detail="File tidak bisa dibaca. Gunakan CSV atau XLSX.")
    df.columns = [str(c).strip().lower() for c in df.columns]
    colmap = {}
    for c in df.columns:
        if "nisn" in c:
            colmap["nisn"] = c
        elif "nis" in c:
            colmap["nis"] = c
        elif "email" in c:
            colmap["parent_email"] = c
        elif "alamat" in c or "address" in c:
            colmap["address"] = c
        elif "hp" in c or "telp" in c or "phone" in c:
            colmap["parent_phone"] = c
        elif "ortu" in c or "wali" in c or "parent" in c:
            colmap["parent_name"] = c
        elif "nama" in c or "name" in c:
            colmap["name"] = c
        elif "kelas" in c or "class" in c or "rombel" in c:
            colmap["class"] = c
        elif "jk" in c or "kelamin" in c or "gender" in c or "l/p" in c:
            colmap["gender"] = c
    if "name" not in colmap:
        raise HTTPException(status_code=400, detail="Kolom 'name'/'nama' wajib ada")
    valid, errors = [], []
    for i, row in df.iterrows():
        def val(key):
            col = colmap.get(key)
            if not col:
                return ""
            v = str(row[col]).strip()
            return "" if v.lower() in ("nan", "none") else v
        nm = val("name")
        if not nm:
            # baris benar-benar kosong (pemisah) -> lewati diam-diam; baris berisi tapi tanpa nama -> tetap dilaporkan
            if any(val(k) for k in ("nis", "nisn", "gender", "class")):
                errors.append({"row": int(i) + 2, "message": "Nama kosong"})
            continue
        valid.append({"name": nm, "nis": val("nis"), "nisn": val("nisn"),
                      "gender": _norm_gender(val("gender")), "class": val("class"),
                      "parent_phone": normalize_phone(val("parent_phone")),
                      "parent_name": val("parent_name"), "parent_email": val("parent_email").lower(),
                      "address": val("address")})
    return {"valid": valid, "errors": errors, "total": len(df)}


class CommitIn(BaseModel):
    rows: List[dict]


@router.post("/admin/students/import/commit")
async def import_commit(body: CommitIn, user: dict = Depends(admin_dep)):
    sid = user["school_id"]
    existing = {s.get("nis") for s in await db.students.find({"school_id": sid}, {"nis": 1}).to_list(10000) if s.get("nis")}
    docs = []
    for r in body.rows:
        nis = str(r.get("nis", "")).strip()
        if nis and nis in existing:
            continue
        docs.append({"id": str(uuid.uuid4()), "school_id": sid, "name": str(r.get("name", "")).strip(),
                     "nis": nis, "nisn": str(r.get("nisn", "")).strip(),
                     "gender": _norm_gender(str(r.get("gender", ""))), "class": str(r.get("class", "")).strip(),
                     "parent_phone": normalize_phone(str(r.get("parent_phone", ""))),
                     "parent_name": str(r.get("parent_name", "")).strip(),
                     "parent_email": str(r.get("parent_email", "")).strip().lower(),
                     "address": str(r.get("address", "")).strip(),
                     "status": "aktif"})
        if nis:
            existing.add(nis)
    if docs:
        await db.students.insert_many(docs)
    return {"inserted": len(docs)}


@router.get("/admin/students/export")
async def export_students(format: str = "xlsx", user: dict = Depends(admin_dep)):
    students = await list_students(user)
    df = pd.DataFrame([{
        "Nama": s.get("name", ""), "NIS": s.get("nis", ""), "NISN": s.get("nisn", ""),
        "L/P": s.get("gender", ""), "Kelas": s.get("class", ""),
        "HP Ortu": s.get("parent_phone", ""), "Nama Ortu": s.get("parent_name", ""),
        "Email Ortu": s.get("parent_email", ""), "Alamat": s.get("address", ""),
        "Enroll Wajah": "Terdaftar" if s.get("enrolled") else "Belum",
    } for s in students])
    buf = io.BytesIO()
    if format == "csv":
        df.to_csv(buf, index=False)
        buf.seek(0)
        return StreamingResponse(buf, media_type="text/csv",
                                 headers={"Content-Disposition": "attachment; filename=siswa.csv"})
    df.to_excel(buf, index=False)
    buf.seek(0)
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=siswa.xlsx"})


# ---------- Leaves ----------
@router.get("/admin/leaves")
async def list_leaves(user: dict = Depends(admin_dep)):
    return await db.leaves.find({"school_id": user["school_id"]}, {"_id": 0}).sort("created_at", -1).to_list(1000)


class DecisionIn(BaseModel):
    status: str  # approved | rejected


@router.post("/admin/leaves/{lid}/decision")
async def decide_leave(lid: str, body: DecisionIn, user: dict = Depends(admin_dep)):
    if body.status not in ("approved", "rejected"):
        raise HTTPException(status_code=400, detail="Status tidak valid")
    res = await db.leaves.update_one(
        {"id": lid, "school_id": user["school_id"]},
        {"$set": {"status": body.status, "decided_by": user["id"], "decided_at": now_iso()}})
    if res.matched_count == 0:
        raise HTTPException(status_code=404, detail="Pengajuan tidak ditemukan")
    return {"ok": True}


class LeaveAdminIn(BaseModel):
    teacher_id: str
    type: str  # izin | sakit | cuti
    date_from: str
    date_to: str
    reason: str = ""


@router.post("/admin/leaves")
async def create_leave_admin(body: LeaveAdminIn, user: dict = Depends(admin_dep)):
    if body.type not in ("izin", "sakit", "cuti"):
        raise HTTPException(status_code=400, detail="Tipe tidak valid")
    t = await db.teachers.find_one({"id": body.teacher_id, "school_id": user["school_id"]}, {"_id": 0, "name": 1})
    if not t:
        raise HTTPException(status_code=404, detail="Guru tidak ditemukan")
    doc = {
        "id": str(uuid.uuid4()), "school_id": user["school_id"], "teacher_id": body.teacher_id,
        "teacher_name": t["name"], "type": body.type, "date_from": body.date_from,
        "date_to": body.date_to, "reason": body.reason, "status": "approved",
        "decided_by": user["id"], "decided_at": now_iso(), "created_at": now_iso(),
    }
    await db.leaves.insert_one(doc)
    doc.pop("_id", None)
    return doc


class LeavePatch(BaseModel):
    type: str | None = None
    date_from: str | None = None
    date_to: str | None = None
    reason: str | None = None


@router.patch("/admin/leaves/{lid}")
async def update_leave(lid: str, body: LeavePatch, user: dict = Depends(admin_dep)):
    upd = body.model_dump(exclude_unset=True)
    if "type" in upd and upd["type"] not in ("izin", "sakit", "cuti"):
        raise HTTPException(status_code=400, detail="Tipe tidak valid")
    if not upd:
        raise HTTPException(status_code=400, detail="Tidak ada perubahan")
    res = await db.leaves.update_one({"id": lid, "school_id": user["school_id"]}, {"$set": upd})
    if res.matched_count == 0:
        raise HTTPException(status_code=404, detail="Pengajuan tidak ditemukan")
    return {"ok": True}


@router.delete("/admin/leaves/{lid}")
async def delete_leave(lid: str, user: dict = Depends(admin_dep)):
    res = await db.leaves.delete_one({"id": lid, "school_id": user["school_id"]})
    if res.deleted_count == 0:
        raise HTTPException(status_code=404, detail="Pengajuan tidak ditemukan")
    return {"ok": True}


# ---------- Absensi Mapel (admin, read-only) ----------
@router.get("/admin/subject-attendance")
async def admin_subject_attendance(date_from: str, date_to: str, class_name: str | None = None,
                                   subject: str | None = None, user: dict = Depends(admin_dep)):
    q = {"school_id": user["school_id"], "date": {"$gte": date_from, "$lte": date_to}}
    if class_name:
        q["class_name"] = class_name
    if subject:
        q["subject"] = subject
    return await db.subject_attendance.find(q, {"_id": 0}).sort(
        [("date", -1), ("class_name", 1), ("subject", 1), ("student_name", 1)]).to_list(10000)


# ---------- Employees (Karyawan) ----------
class EmployeeIn(BaseModel):
    name: str
    email: EmailStr
    password: str
    nip: str = ""
    department: str = ""
    position: str = ""
    overtime_rate: int | None = None
    base_salary: int | None = None
    card_uid: str = ""
    shift_id: str = ""


class EmployeePatch(BaseModel):
    name: str | None = None
    nip: str | None = None
    department: str | None = None
    position: str | None = None
    active: bool | None = None
    overtime_rate: int | None = None
    base_salary: int | None = None
    card_uid: str | None = None
    shift_id: str | None = None


@router.get("/admin/employees")
async def list_employees(user: dict = Depends(admin_dep)):
    emps = await db.employees.find({"school_id": user["school_id"]}, {"_id": 0, "embedding": 0, "photo": 0}).to_list(2000)
    enrolled_ids = {e["id"] for e in await db.employees.find(
        {"school_id": user["school_id"], "embedding": {"$type": "array"}}, {"_id": 0, "id": 1}).to_list(2000)}
    uids = [e["user_id"] for e in emps if e.get("user_id")]
    users = {u["id"]: u["email"] for u in await db.users.find({"id": {"$in": uids}}, {"_id": 0, "id": 1, "email": 1}).to_list(2000)}
    for e in emps:
        e["email"] = users.get(e.get("user_id"), "")
        e["enrolled"] = e["id"] in enrolled_ids
    return emps


@router.post("/admin/employees")
async def create_employee(body: EmployeeIn, user: dict = Depends(admin_dep)):
    if await db.users.find_one({"email": body.email.lower()}):
        raise HTTPException(status_code=400, detail="Email sudah dipakai")
    uid = str(uuid.uuid4())
    await db.users.insert_one({
        "id": uid, "email": body.email.lower(), "name": body.name, "role": "employee",
        "password_hash": hash_password(body.password), "school_id": user["school_id"], "created_at": now_iso(),
    })
    emp = {
        "id": str(uuid.uuid4()), "school_id": user["school_id"], "user_id": uid,
        "name": body.name, "nip": body.nip, "department": body.department, "position": body.position,
        "overtime_rate": body.overtime_rate, "base_salary": body.base_salary, "card_uid": body.card_uid.strip(),
        "shift_id": body.shift_id,
        "embedding": None, "photo": None, "active": True, "created_at": now_iso(),
    }
    await db.employees.insert_one(emp)
    emp.pop("_id", None)
    emp.pop("embedding", None)
    emp.pop("photo", None)
    emp["email"] = body.email.lower()
    emp["enrolled"] = False
    return emp


@router.patch("/admin/employees/{eid}")
async def update_employee(eid: str, body: EmployeePatch, user: dict = Depends(admin_dep)):
    upd = {k: v for k, v in body.model_dump().items() if v is not None}
    update = {}
    if upd:
        update["$set"] = upd
    unset = {f: "" for f in ("overtime_rate", "base_salary") if f in body.model_fields_set and getattr(body, f) is None}
    if unset:
        update["$unset"] = unset
    if not update:
        raise HTTPException(status_code=400, detail="Tidak ada perubahan")
    await db.employees.update_one({"id": eid, "school_id": user["school_id"]}, update)
    if "name" in upd:
        e = await db.employees.find_one({"id": eid}, {"_id": 0, "user_id": 1})
        if e:
            await db.users.update_one({"id": e["user_id"]}, {"$set": {"name": upd["name"]}})
    return {"ok": True}


@router.delete("/admin/employees/{eid}")
async def delete_employee(eid: str, user: dict = Depends(admin_dep)):
    e = await db.employees.find_one({"id": eid, "school_id": user["school_id"]})
    if not e:
        raise HTTPException(status_code=404, detail="Karyawan tidak ditemukan")
    await db.employees.delete_one({"id": eid})
    await db.users.delete_one({"id": e["user_id"]})
    return {"ok": True}


@router.post("/admin/employees/{eid}/enroll")
async def enroll_employee_face(eid: str, body: EnrollIn, user: dict = Depends(admin_dep)):
    e = await db.employees.find_one({"id": eid, "school_id": user["school_id"]})
    if not e:
        raise HTTPException(status_code=404, detail="Karyawan tidak ditemukan")
    try:
        emb = await run_in_threadpool(embed, body.photo)
    except NoFaceError as exc:
        raise HTTPException(status_code=422, detail=str(exc))
    except Exception:
        raise HTTPException(status_code=400, detail="Foto tidak valid")
    await _face_dup_check(user["school_id"], emb, eid)
    await db.employees.update_one({"id": eid}, {"$set": {"embedding": emb, "embedding_model": "buffalo_s", "photo": body.photo, "enrolled_at": now_iso()}})
    return {"ok": True, "enrolled": True}


# ---------- Overtime (Lembur) ----------
@router.get("/admin/overtime")
async def list_overtime(status: str | None = None, user: dict = Depends(admin_dep)):
    q = {"school_id": user["school_id"]}
    if status in ("pending", "approved", "rejected"):
        q["status"] = status
    return await db.overtime_requests.find(q, {"_id": 0}).sort("created_at", -1).to_list(500)


@router.post("/admin/overtime/{oid}/decision")
async def decide_overtime(oid: str, body: DecisionIn, user: dict = Depends(admin_dep)):
    if body.status not in ("approved", "rejected"):
        raise HTTPException(status_code=400, detail="Status tidak valid")
    res = await db.overtime_requests.update_one(
        {"id": oid, "school_id": user["school_id"], "status": "pending"},
        {"$set": {"status": body.status, "decided_by": user["id"], "decided_at": now_iso()}})
    if res.matched_count == 0:
        raise HTTPException(status_code=404, detail="Pengajuan tidak ditemukan")
    return {"ok": True}


async def _overtime_recap(sid: str, date_from: str, date_to: str) -> list:
    settings = await db.settings.find_one({"school_id": sid}, {"_id": 0, "overtime_rate": 1}) or {}
    default_rate = int(settings.get("overtime_rate") or 0)
    emps = await db.employees.find({"school_id": sid}, {"_id": 0, "embedding": 0, "photo": 0}).to_list(2000)
    outs = await db.attendance.find(
        {"school_id": sid, "person_type": "employee", "type": "out",
         "date": {"$gte": date_from, "$lte": date_to}, "overtime_minutes": {"$gt": 0}},
        {"_id": 0, "employee_id": 1, "date": 1, "overtime_minutes": 1}).to_list(20000)
    actual = {}
    for r in outs:
        k = (r.get("employee_id"), r["date"])
        actual[k] = actual.get(k, 0) + r.get("overtime_minutes", 0)
    reqs = await db.overtime_requests.find(
        {"school_id": sid, "status": "approved", "date": {"$gte": date_from, "$lte": date_to}},
        {"_id": 0, "employee_id": 1, "date": 1, "minutes": 1}).to_list(20000)
    approved = {}
    for r in reqs:
        k = (r["employee_id"], r["date"])
        approved[k] = approved.get(k, 0) + r.get("minutes", 0)
    recap = []
    for e in emps:
        rate = e.get("overtime_rate") if e.get("overtime_rate") is not None else default_rate
        tot_actual = tot_paid = 0
        for (eid, d), mins in actual.items():
            if eid != e["id"]:
                continue
            tot_actual += mins
            tot_paid += min(mins, approved.get((eid, d), 0))
        recap.append({
            "id": e["id"], "name": e["name"], "nip": e.get("nip", ""),
            "department": e.get("department", ""), "position": e.get("position", ""),
            "overtime_rate": rate, "overtime_minutes": tot_actual,
            "paid_minutes": tot_paid, "overtime_pay": round(tot_paid / 60 * rate),
            "base_salary": e.get("base_salary"),
        })
    recap.sort(key=lambda r: r["name"])
    return recap


@router.get("/admin/reports/overtime")
async def report_overtime(date_from: str, date_to: str, user: dict = Depends(admin_dep)):
    return await _overtime_recap(user["school_id"], date_from, date_to)


@router.get("/admin/reports/overtime/export")
async def report_overtime_export(date_from: str, date_to: str, user: dict = Depends(admin_dep)):
    recap = await _overtime_recap(user["school_id"], date_from, date_to)
    df = pd.DataFrame([{
        "Nama": r["name"], "NIP": r["nip"], "Departemen": r["department"], "Jabatan": r["position"],
        "Lembur Aktual (mnt)": r["overtime_minutes"], "Lembur Disetujui (mnt)": r["paid_minutes"],
        "Tarif/Jam (Rp)": r["overtime_rate"], "Upah Lembur (Rp)": r["overtime_pay"],
    } for r in recap])
    buf = io.BytesIO()
    df.to_excel(buf, index=False)
    buf.seek(0)
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=laporan-lembur.xlsx"})


# ---------- Payroll (Penggajian) ----------
async def _payroll_rows(sid: str, period: str) -> list:
    y, m = int(period[:4]), int(period[5:7])
    date_from = f"{period}-01"
    date_to = f"{period}-{calendar.monthrange(y, m)[1]:02d}"
    recap = await _overtime_recap(sid, date_from, date_to)
    ins = await db.attendance.find(
        {"school_id": sid, "person_type": "employee", "type": "in",
         "date": {"$gte": date_from, "$lte": date_to}},
        {"_id": 0, "employee_id": 1, "date": 1}).to_list(20000)
    days = {}
    for r in ins:
        days.setdefault(r.get("employee_id"), set()).add(r["date"])
    rows = []
    for r in recap:
        base = r.get("base_salary") or 0
        rows.append({**r, "base_salary": base, "present_days": len(days.get(r["id"], set())),
                     "total_pay": base + r["overtime_pay"]})
    return rows


def _check_period(period: str):
    if not re.match(r"^(19|20)\d{2}-(0[1-9]|1[0-2])$", period):
        raise HTTPException(status_code=422, detail="Periode tidak valid")


@router.get("/admin/reports/payroll")
async def report_payroll(period: str, user: dict = Depends(admin_dep)):
    _check_period(period)
    return await _payroll_rows(user["school_id"], period)


@router.get("/admin/reports/payroll/export")
async def report_payroll_export(period: str, user: dict = Depends(admin_dep)):
    _check_period(period)
    rows = await _payroll_rows(user["school_id"], period)
    df = pd.DataFrame([{
        "Nama": r["name"], "NIP": r["nip"], "Departemen": r["department"], "Jabatan": r["position"],
        "Hadir (hari)": r["present_days"], "Lembur Disetujui (mnt)": r["paid_minutes"],
        "Tarif Lembur/Jam (Rp)": r["overtime_rate"], "Upah Lembur (Rp)": r["overtime_pay"],
        "Gaji Pokok (Rp)": r["base_salary"], "Total (Rp)": r["total_pay"],
    } for r in rows])
    buf = io.BytesIO()
    df.to_excel(buf, index=False)
    buf.seek(0)
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename=penggajian-{period}.xlsx"})



# ---------- Sesi absen mapel (kunci/buka oleh admin) ----------
@router.get("/admin/subject-sessions")
async def subject_sessions(date: str, user: dict = Depends(admin_dep)):
    pipeline = [
        {"$match": {"school_id": user["school_id"], "date": date}},
        {"$group": {"_id": {"teacher_id": "$teacher_id", "subject": "$subject", "class_name": "$class_name"},
                    "teacher_name": {"$first": "$teacher_name"},
                    "count": {"$sum": 1},
                    "locked": {"$max": {"$cond": [{"$eq": ["$locked", True]}, 1, 0]}}}},
        {"$sort": {"_id.class_name": 1, "_id.subject": 1}},
    ]
    rows = await db.subject_attendance.aggregate(pipeline).to_list(500)
    return [{"teacher_id": r["_id"]["teacher_id"], "subject": r["_id"]["subject"],
             "class_name": r["_id"]["class_name"], "teacher_name": r["teacher_name"],
             "count": r["count"], "locked": bool(r["locked"])} for r in rows]


class SessionLockIn(BaseModel):
    date: str
    teacher_id: str
    subject: str
    class_name: str
    lock: bool


@router.post("/admin/subject-sessions/lock")
async def lock_subject_session(body: SessionLockIn, user: dict = Depends(admin_dep)):
    key = {"school_id": user["school_id"], "date": body.date, "teacher_id": body.teacher_id,
           "subject": body.subject, "class_name": body.class_name}
    if body.lock:
        r = await db.subject_attendance.update_many(key, {"$set": {"locked": True, "locked_by": user["email"]}})
    else:
        r = await db.subject_attendance.update_many(key, {"$unset": {"locked": "", "locked_by": ""}})
    return {"ok": True, "modified": r.modified_count}


# ---------- Reports ----------
@router.get("/admin/reports/attendance")
async def report_attendance(date_from: str, date_to: str, teacher_id: str | None = None,
                            person: str | None = None, user: dict = Depends(admin_dep)):
    q = {"school_id": user["school_id"], "date": {"$gte": date_from, "$lte": date_to}}
    if teacher_id:
        q["teacher_id"] = teacher_id
    if person in ("student", "teacher", "employee"):
        q["person_type"] = person
    rows = await db.attendance.find(q, {"_id": 0, "photo": 0}).sort([("date", 1), ("ts_server", 1)]).to_list(10000)
    students = {s["id"]: s for s in await db.students.find(
        {"school_id": user["school_id"]}, {"_id": 0, "id": 1, "nisn": 1, "gender": 1}).to_list(10000)}
    for r in rows:
        r["time"] = r.get("time_local") or r.get("ts_device", r.get("ts_server", ""))[11:16]
        st = students.get(r.get("student_id"))
        if st:
            r["nisn"] = st.get("nisn", "")
            r["gender"] = st.get("gender", "")
    return rows


@router.get("/admin/reports/export")
async def report_export(format: str, date_from: str, date_to: str, person: str | None = None,
                        user: dict = Depends(admin_dep)):
    rows = await report_attendance(date_from, date_to, None, person, user)
    if format == "xlsx":
        df = pd.DataFrame([{
            "Tanggal": r.get("date"), "Nama": r.get("teacher_name"), "NISN": r.get("nisn", ""),
            "L/P": r.get("gender", ""), "Tipe": r.get("type"),
            "Jam": r.get("time"), "Status": r.get("status"), "Telat (mnt)": r.get("late_minutes", 0),
            "Lembur (mnt)": r.get("overtime_minutes", 0), "Offline": "Ya" if r.get("offline") else "Tidak",
        } for r in rows])
        buf = io.BytesIO()
        df.to_excel(buf, index=False)
        buf.seek(0)
        return StreamingResponse(
            buf,
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={"Content-Disposition": "attachment; filename=laporan-absensi.xlsx"})
    school = await db.schools.find_one({"id": user["school_id"]}, {"_id": 0, "name": 1})
    path = build_report_pdf(f"/tmp/report_{user['school_id']}.pdf", school["name"], date_from, date_to, rows)
    return FileResponse(path, media_type="application/pdf", filename="laporan-absensi.pdf")



# ---------- Laporan kehadiran per orang: siswa / guru / karyawan (admin) ----------
PEOPLE_REPORT_CFG = {
    "student": {"coll": "students", "idf": "student_id", "extra": {"status": {"$ne": "lulus"}},
                "ref": "nis", "grp": "class", "ref_label": "NIS", "grp_label": "Kelas",
                "label": "Siswa", "slug": "siswa"},
    "teacher": {"coll": "teachers", "idf": "teacher_id", "extra": {"active": True},
                "ref": "nip", "grp": "subject", "ref_label": "NIP", "grp_label": "Mapel",
                "label": "Guru", "slug": "guru"},
    "employee": {"coll": "employees", "idf": "employee_id", "extra": {"active": True},
                 "ref": "nip", "grp": "department", "ref_label": "NIP", "grp_label": "Departemen",
                 "label": "Karyawan", "slug": "karyawan"},
}


def _person_cfg(person: str) -> dict:
    cfg = PEOPLE_REPORT_CFG.get(person)
    if not cfg:
        raise HTTPException(status_code=422, detail="Tipe person tidak valid")
    return cfg


async def _people_report_rows(sid: str, person: str, date_from: str, date_to: str, class_name: str | None):
    cfg = _person_cfg(person)
    q = {"school_id": sid, "person_type": person, "date": {"$gte": date_from, "$lte": date_to}}
    if class_name and person == "student":
        q["class"] = class_name
    rows = await db.attendance.find(q, {"_id": 0, "photo": 0}).sort([("date", 1), ("ts_server", 1)]).to_list(10000)
    for r in rows:
        r["time"] = r.get("time_local") or r.get("ts_device", r.get("ts_server", ""))[11:16]
    pq = {"school_id": sid, **cfg["extra"]}
    if class_name and person == "student":
        pq["class"] = class_name
    persons = await db[cfg["coll"]].find(
        pq, {"_id": 0, "id": 1, "name": 1, "nis": 1, "nip": 1, "class": 1, "subject": 1, "department": 1}).to_list(5000)
    return rows, persons, cfg


def _people_recap(rows: list, persons: list, cfg: dict) -> list:
    active_days = {r["date"] for r in rows}
    idf = cfg["idf"]
    recap = []
    for p in persons:
        rs = [r for r in rows if r.get(idf) == p["id"]]
        hadir = {r["date"] for r in rs if r.get("att_status", "present") == "present"}
        telat = {r["date"] for r in rs if r.get("att_status", "present") == "present" and r.get("late_minutes", 0) > 0}
        sakit = {r["date"] for r in rs if r.get("att_status") == "sakit"}
        izin = {r["date"] for r in rs if r.get("att_status") == "izin"}
        alpha = max(0, len(active_days) - len(hadir | sakit | izin))
        recap.append({"id": p["id"], "name": p["name"], "ref": p.get(cfg["ref"], "") or "",
                      "group": p.get(cfg["grp"], "") or "",
                      "hadir": len(hadir), "telat": len(telat), "sakit": len(sakit), "izin": len(izin),
                      "alpha": alpha, "active_days": len(active_days)})
    recap.sort(key=lambda r: (r["group"], r["name"]))
    return recap


@router.get("/admin/reports/people")
async def report_people_daily(person: str, date_from: str, date_to: str, class_name: str | None = None,
                              user: dict = Depends(admin_dep)):
    rows, _, _ = await _people_report_rows(user["school_id"], person, date_from, date_to, class_name)
    return rows


@router.get("/admin/reports/people/recap")
async def report_people_recap(person: str, date_from: str, date_to: str, class_name: str | None = None,
                              user: dict = Depends(admin_dep)):
    rows, persons, cfg = await _people_report_rows(user["school_id"], person, date_from, date_to, class_name)
    return _people_recap(rows, persons, cfg)


@router.get("/admin/reports/people/export")
async def report_people_export(person: str, format: str, kind: str, date_from: str, date_to: str,
                               class_name: str | None = None, user: dict = Depends(admin_dep)):
    rows, persons, cfg = await _people_report_rows(user["school_id"], person, date_from, date_to, class_name)
    school = await db.schools.find_one({"id": user["school_id"]}, {"_id": 0, "name": 1})
    xlsx_mime = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    if kind == "recap":
        data = _people_recap(rows, persons, cfg)
        if format == "xlsx":
            df = pd.DataFrame([{
                "Nama": r["name"], cfg["ref_label"]: r["ref"], cfg["grp_label"]: r["group"],
                "Hadir": r["hadir"], "Telat": r["telat"], "Sakit": r["sakit"], "Izin": r["izin"],
                "Alpha": r["alpha"], "Hari Efektif": r["active_days"],
            } for r in data])
            buf = io.BytesIO()
            df.to_excel(buf, index=False)
            buf.seek(0)
            return StreamingResponse(
                buf, media_type=xlsx_mime,
                headers={"Content-Disposition": f"attachment; filename=rekap-kehadiran-{cfg['slug']}.xlsx"})
        path = build_recap_pdf(f"/tmp/recap_{person}_{user['school_id']}.pdf", school["name"],
                               date_from, date_to, data, class_name, cfg["label"], cfg["grp_label"])
        return FileResponse(path, media_type="application/pdf",
                            filename=f"rekap-kehadiran-{cfg['slug']}.pdf")
    pmap = {p["id"]: p for p in persons}
    idf = cfg["idf"]
    if format == "xlsx":
        df = pd.DataFrame([{
            "Tanggal": r.get("date"), "Nama": r.get("teacher_name"),
            cfg["ref_label"]: (pmap.get(r.get(idf)) or {}).get(cfg["ref"], "") or "",
            cfg["grp_label"]: r.get("class", "") or r.get("department", "") or "",
            "Tipe": "Masuk" if r.get("type") == "in" else "Pulang",
            "Jam": r.get("time"),
            "Status": (r.get("att_status") if r.get("att_status") not in (None, "present") else r.get("status")),
            "Telat (mnt)": r.get("late_minutes", 0), "Lembur (mnt)": r.get("overtime_minutes", 0),
            "Offline": "Ya" if r.get("offline") else "Tidak",
        } for r in rows])
        buf = io.BytesIO()
        df.to_excel(buf, index=False)
        buf.seek(0)
        return StreamingResponse(
            buf, media_type=xlsx_mime,
            headers={"Content-Disposition": f"attachment; filename=laporan-{cfg['slug']}.xlsx"})
    path = build_report_pdf(f"/tmp/report_{person}_{user['school_id']}.pdf", school["name"], date_from, date_to, rows)
    return FileResponse(path, media_type="application/pdf", filename=f"laporan-{cfg['slug']}.pdf")